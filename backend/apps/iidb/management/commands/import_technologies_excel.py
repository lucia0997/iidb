import tkinter as tk
from tkinter import filedialog
from decimal import Decimal, InvalidOperation
from django.core.management.base import BaseCommand
from django.db import transaction
from apps.iidb.models import Technology, TRL, TechnologyTRL


class Command(BaseCommand):
    help = "Importa tecnologías desde un archivo Excel seleccionado mediante un explorador de archivos"

    def handle(self, *args, **options):
        # Abrir explorador de archivos
        root = tk.Tk()
        root.withdraw()  # Ocultar la ventana principal
        
        file_path = filedialog.askopenfilename(
            title="Seleccionar archivo Excel",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")]
        )
        
        root.destroy()
        
        if not file_path:
            self.stdout.write(self.style.WARNING("No se seleccionó ningún archivo. Operación cancelada."))
            return
        
        self.stdout.write(f"Leyendo archivo: {file_path}")
        
        try:
            import openpyxl
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            sheet = workbook.active
            
            # Leer encabezados de la primera fila
            headers = {}
            for col_idx, cell in enumerate(sheet[1], start=1):
                if cell.value:
                    headers[cell.value] = col_idx
            
            self.stdout.write(f"Encabezados encontrados: {list(headers.keys())}")
            
            # Mapeo de columnas del Excel a campos de la base de datos
            column_mapping = {
                'Physical Technology Cluster': 'physical_technology_cluster',
                'Digital Technology Cluster': 'digital_technology_cluster',
                'Product Roadmap': 'product_roadmap',
                'Technology Roadmap': 'technology_roadmap',
                'Technology Name': 'technology_name',
                'Technology Description': 'technology_description',
                'Current TRL (1-9)': 'current_trl',
                'Dependencies other Techno': 'dependencies',
                'FoM Type': 'fom_type',
                'FoM Value (%)': 'fom_value_percent',
                'Targeted Programmes A/C Application': 'targeted_programmes_ac',
            }
            
            # Columnas de años TRL
            trl_year_columns = {}
            for i in range(1, 10):
                col_name = f'TRL{i} Year'
                if col_name in headers:
                    trl_year_columns[i] = headers[col_name]
            
            # Columnas de costos TRL
            trl_cost_columns = {}
            for i in range(1, 10):
                col_name = f'TRL{i} Cost (k€)'
                if col_name in headers:
                    trl_cost_columns[i] = headers[col_name]
            
            technologies_created = 0
            technologies_updated = 0
            errors = []
            
            # Procesar cada fila de datos (empezando desde la fila 2)
            for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=False), start=2):
                # Verificar si la fila tiene datos (al menos el nombre de la tecnología)
                tech_name_col = headers.get('Technology Name') or headers.get('Technology Nam')
                if not tech_name_col:
                    continue
                
                tech_name_cell = sheet.cell(row=row_idx, column=tech_name_col)
                if not tech_name_cell.value or str(tech_name_cell.value).strip() == '':
                    continue
                
                try:
                    with transaction.atomic():
                        # Extraer datos básicos
                        tech_data = {
                            'targeted_programmes': [],  # Inicializar por defecto
                            'ac_application': '',  # Inicializar por defecto
                        }
                        for excel_col, db_field in column_mapping.items():
                            if excel_col in headers:
                                col_idx = headers[excel_col]
                                cell_value = sheet.cell(row=row_idx, column=col_idx).value
                                
                                if cell_value is not None:
                                    if db_field == 'current_trl':
                                        # Convertir a entero
                                        try:
                                            tech_data[db_field] = int(cell_value) if str(cell_value).strip() not in ['N/A', '', None] else None
                                        except (ValueError, TypeError):
                                            tech_data[db_field] = None
                                    elif db_field == 'fom_value_percent':
                                        # Convertir a Decimal
                                        try:
                                            # Limpiar el valor (puede tener comas o comillas)
                                            value_str = str(cell_value).replace(',', '.').strip('"').strip()
                                            tech_data[db_field] = Decimal(value_str) if value_str and value_str != 'N/A' else None
                                        except (InvalidOperation, ValueError, TypeError):
                                            tech_data[db_field] = None
                                    elif db_field == 'dependencies':
                                        # Convertir a lista JSON
                                        if str(cell_value).strip() in ['N/A', '', 'None', None]:
                                            tech_data[db_field] = []
                                        else:
                                            # Si es una cadena, convertir a lista
                                            deps = str(cell_value).strip()
                                            tech_data[db_field] = [deps] if deps else []
                                    elif db_field == 'targeted_programmes_ac':
                                        # Separar targeted_programmes y ac_application
                                        value_str = str(cell_value).strip()
                                        if value_str and value_str != 'N/A':
                                            # Separar por comas
                                            parts = [p.strip() for p in value_str.split(',')]
                                            
                                            # Intentar identificar programas vs aplicaciones
                                            # Basado en el formato: "FCAS RC, FCAS NGF, Fuselage, Wings, Tail"
                                            # Los programas suelen ser más cortos y específicos
                                            programmes = []
                                            ac_parts = []
                                            
                                            # Heurística simple: si contiene "Fuselage", "Wings", "Tail" son aplicaciones
                                            application_keywords = ['Fuselage', 'Wings', 'Tail', 'A/C', 'Application']
                                            
                                            for part in parts:
                                                if any(keyword in part for keyword in application_keywords):
                                                    ac_parts.append(part)
                                                else:
                                                    programmes.append(part)
                                            
                                            # Si no se encontraron aplicaciones, usar todo como ac_application
                                            if not ac_parts and programmes:
                                                tech_data['ac_application'] = value_str
                                                tech_data['targeted_programmes'] = []
                                            else:
                                                tech_data['ac_application'] = ', '.join(ac_parts) if ac_parts else ''
                                                tech_data['targeted_programmes'] = programmes if programmes else []
                                        else:
                                            tech_data['ac_application'] = ''
                                            tech_data['targeted_programmes'] = []
                                    else:
                                        # Campo de texto normal
                                        value_str = str(cell_value).strip()
                                        tech_data[db_field] = value_str if value_str != 'N/A' else ''
                                
                        # Crear o actualizar la tecnología
                        technology_name = tech_data.get('technology_name', '').strip()
                        if not technology_name:
                            errors.append(f"Fila {row_idx}: No se encontró nombre de tecnología")
                            continue
                        
                        technology, created = Technology.objects.update_or_create(
                            technology_name=technology_name,
                            defaults={
                                'physical_technology_cluster': tech_data.get('physical_technology_cluster', ''),
                                'digital_technology_cluster': tech_data.get('digital_technology_cluster', ''),
                                'product_roadmap': tech_data.get('product_roadmap', ''),
                                'technology_roadmap': tech_data.get('technology_roadmap', ''),
                                'technology_description': tech_data.get('technology_description', ''),
                                'current_trl': tech_data.get('current_trl'),
                                'dependencies': tech_data.get('dependencies', []),
                                'fom_type': tech_data.get('fom_type', ''),
                                'fom_value_percent': tech_data.get('fom_value_percent'),
                                'targeted_programmes': tech_data.get('targeted_programmes', []),
                                'ac_application': tech_data.get('ac_application', ''),
                            }
                        )
                        
                        if created:
                            technologies_created += 1
                        else:
                            technologies_updated += 1
                        
                        # Procesar TRLs
                        for trl_num in range(1, 10):
                            trl_year = None
                            trl_cost = None
                            
                            # Obtener año TRL
                            if trl_num in trl_year_columns:
                                year_col = trl_year_columns[trl_num]
                                year_value = sheet.cell(row=row_idx, column=year_col).value
                                if year_value and str(year_value).strip() not in ['N/A', '', None]:
                                    try:
                                        trl_year = int(year_value)
                                    except (ValueError, TypeError):
                                        pass
                            
                            # Obtener costo TRL
                            if trl_num in trl_cost_columns:
                                cost_col = trl_cost_columns[trl_num]
                                cost_value = sheet.cell(row=row_idx, column=cost_col).value
                                if cost_value and str(cost_value).strip() not in ['N/A', '', None]:
                                    try:
                                        # Convertir de k€ a €
                                        cost_decimal = Decimal(str(cost_value).replace(',', '.'))
                                        trl_cost = cost_decimal * 1000  # Convertir k€ a €
                                    except (InvalidOperation, ValueError, TypeError):
                                        pass
                            
                            # Crear TRL si hay año o costo
                            if trl_year or trl_cost:
                                trl, _ = TRL.objects.get_or_create(
                                    trl_number=trl_num,
                                    trl_year=trl_year,
                                    trl_cost=trl_cost,
                                    defaults={
                                        'trl_number': trl_num,
                                        'trl_year': trl_year,
                                        'trl_cost': trl_cost,
                                    }
                                )
                                
                                # Relacionar tecnología con TRL
                                TechnologyTRL.objects.get_or_create(
                                    technology=technology,
                                    trl=trl,
                                    defaults={'trl_number': trl_num}
                                )
                
                except Exception as e:
                    error_msg = f"Fila {row_idx}: Error al procesar - {str(e)}"
                    errors.append(error_msg)
                    self.stdout.write(self.style.ERROR(error_msg))
            
            # Mostrar resumen
            self.stdout.write(self.style.SUCCESS(f"\n✓ Tecnologías creadas: {technologies_created}"))
            self.stdout.write(self.style.SUCCESS(f"✓ Tecnologías actualizadas: {technologies_updated}"))
            
            if errors:
                self.stdout.write(self.style.WARNING(f"\n⚠ Errores encontrados: {len(errors)}"))
                for error in errors[:10]:  # Mostrar solo los primeros 10 errores
                    self.stdout.write(self.style.ERROR(f"  - {error}"))
                if len(errors) > 10:
                    self.stdout.write(self.style.WARNING(f"  ... y {len(errors) - 10} errores más"))
            else:
                self.stdout.write(self.style.SUCCESS("✓ Proceso completado sin errores"))
        
        except Exception as e:
            self.stdout.write(self.style.ERROR(f"Error al procesar el archivo: {str(e)}"))
            import traceback
            self.stdout.write(self.style.ERROR(traceback.format_exc()))

